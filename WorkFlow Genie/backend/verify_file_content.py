"""
Utility to verify file content and show detailed information
Run this to see what's actually in the file on disk
"""

from pathlib import Path
import hashlib
from datetime import datetime
from loguru import logger
import openpyxl

def verify_file_content(file_path: str):
    """Show detailed file information to verify updates"""
    
    path = Path(file_path).resolve()
    
    logger.info("="*70)
    logger.info("FILE CONTENT VERIFICATION")
    logger.info("="*70)
    
    if not path.exists():
        logger.error(f"File does not exist: {path}")
        return
    
    # Get file stats
    stat = path.stat()
    size = stat.st_size
    mtime = datetime.fromtimestamp(stat.st_mtime)
    
    logger.info(f"File: {path.name}")
    logger.info(f"Full Path: {path}")
    logger.info(f"Size: {size:,} bytes")
    logger.info(f"Last Modified: {mtime.strftime('%Y-%m-%d %H:%M:%S')}")
    
    # Read file content
    with open(path, 'rb') as f:
        content = f.read()
    
    # Calculate hash
    file_hash = hashlib.sha256(content).hexdigest()
    logger.info(f"SHA256 Hash: {file_hash[:32]}...")
    logger.info(f"Short Hash: {file_hash[:16]}")
    
    # Try to open as Excel
    logger.info("\n" + "="*70)
    logger.info("EXCEL CONTENT PREVIEW")
    logger.info("="*70)
    
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        logger.success(f"Excel file opened successfully")
        logger.info(f"Number of sheets: {len(wb.sheetnames)}")
        logger.info(f"Sheet names: {', '.join(wb.sheetnames)}")
        
        # Show first sheet preview
        if wb.sheetnames:
            sheet = wb[wb.sheetnames[0]]
            max_row = sheet.max_row
            max_col = sheet.max_column
            
            logger.info(f"\nFirst sheet: '{wb.sheetnames[0]}'")
            logger.info(f"Dimensions: {max_row} rows x {max_col} columns")
            
            # Show first few rows
            logger.info("\nFirst 5 rows:")
            for row_idx, row in enumerate(sheet.iter_rows(max_row=5, values_only=True), 1):
                row_str = " | ".join([str(cell) if cell is not None else "" for cell in row[:10]])
                if len(row) > 10:
                    row_str += " ..."
                logger.info(f"  Row {row_idx}: {row_str}")
        
        wb.close()
        
    except Exception as e:
        logger.error(f"Could not read Excel content: {e}")
    
    logger.info("\n" + "="*70)
    logger.success("VERIFICATION COMPLETE")
    logger.info("="*70)
    logger.info("\nIf you have this file open in Excel:")
    logger.info("1. Close the file in Excel")
    logger.info("2. Reopen it")
    logger.info("3. Excel caches file contents and won't show updates until reopened")
    logger.info("\nOr check the 'Last Modified' timestamp above - it should be recent!")


if __name__ == "__main__":
    import sys
    
    if len(sys.argv) < 2:
        print("Usage: python verify_file_content.py <file_path>")
        print("\nExample:")
        print('python verify_file_content.py "data\\excel_files\\100f49aa-3ad9-4115-b209-503f9da44c89_TrialBal.xlsx"')
        sys.exit(1)
    
    file_path = sys.argv[1]
    verify_file_content(file_path)
